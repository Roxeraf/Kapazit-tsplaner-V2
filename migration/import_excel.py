"""
Einmal-Import bestehender Excel-Planungsdaten (.xlsm) in den Kapazitätsplaner.
Siehe CONCEPT.md, Abschnitt 8.

Liest jedes Projektblatt (alles außer Konfiguration/Legende/Gesamtuebersicht)
nach der Struktur, die das VBA-Tool erzeugt (legacy/VBA.txt, Funktion
`Exportiere_Als_JSON_Und_PPTX`):

  Zeile 1, Spalte A         Projektname
  Zeile 2, Spalte B..       Monatsspalten
  Teilprojekt 1: Header Zeile 4, Phasenzeilen 5-9  (p,k,t,g,?)
  Teilprojekt 2: Header Zeile 10, Phasenzeilen 11-15
  Teilprojekt 3: Header Zeile 16, Phasenzeilen 17-21
  FTE (Soll):    Zeile 24 (TP1), 25 (TP2), 26 (TP3) — Zeile 27 (Summe) wird ignoriert,
                 die Summe errechnet die neue Web-App aus TP1-3 von selbst.

Das Konfigurationsblatt liefert Startmonat/Anzahl Monate (Row 17/18, Spalte C).

Seit dem Legacy Cutover (Phase 26.9, CONCEPT.md Abschnitt 11 Punkt 26) schreibt der
Import nicht mehr in das alte Gantt/FTE-Grid, sondern in die führende Zielarchitektur:
  - Phasenzellen -> PlanPhase (aufeinanderfolgende Monate mit gleichem Phasencode werden
    zu EINER Phase mit forecast_start/forecast_end verdichtet; "?"-Zellen -> Milestone)
  - FTE-Soll -> ResourceDemand je Monat (Rolle "Allgemein", wird einmalig angelegt)
Grund: Das alte Grid ist entfernt, PUT /projects/.../phasen|fte existiert nicht mehr.

Nutzung:
    python migration/import_excel.py <pfad-zur-datei.xlsm> [--api http://localhost:8000]

Voraussetzung: `pip install openpyxl requests` (nicht Teil von backend/requirements.txt,
da nur für den einmaligen Migrationslauf benötigt).
"""

import argparse
import datetime
import sys

import requests
from openpyxl import load_workbook

PHASE_ROW_CODES = ["p", "k", "t", "g", "?"]
TP_START_ROWS = [5, 11, 17]  # erste Phasenzeile je Teilprojekt
TP_FTE_ROWS = [24, 25, 26]
MONAT_HEADER_ROW = 2
FIRST_MONAT_COL = 2  # Spalte B

SKIP_SHEETS = {"Konfiguration", "Legende", "Gesamtuebersicht"}

# Phasencode -> sprechendes Label, identisch zu backend/app/constants.py PHASE_LABELS
# (Import-Ziel: PlanPhase.phase_type ist Freitext; dieselben Labels, die PlanPhaseList
# als Datalist vorschlägt und die der PPTX-Export wieder auf Kürzel zurückführt).
PHASE_LABELS = {
    "p": "Pflichtenheft",
    "k": "Konfiguration",
    "t": "Test",
    "s": "Schulung",
    "g": "GoLive",
    "?": "Meilenstein",
}

MONAT_NAMEN = [
    "Jan", "Feb", "Mrz", "Apr", "Mai", "Jun",
    "Jul", "Aug", "Sep", "Okt", "Nov", "Dez",
]


def _monat_start(period: str) -> str:
    name, jahr_str = period.split()
    jahr = 2000 + int(jahr_str)
    monat = MONAT_NAMEN.index(name) + 1
    return f"{jahr}-{monat:02d}-01"


def _monat_end(period: str) -> str:
    name, jahr_str = period.split()
    jahr = 2000 + int(jahr_str)
    monat = MONAT_NAMEN.index(name) + 1
    if monat == 12:
        return f"{jahr}-12-31"
    return (datetime.date(jahr, monat + 1, 1) - datetime.timedelta(days=1)).isoformat()


def _berechne_monate(start_monat: str, anzahl_monate: int) -> list[str]:
    """start_monat 'MM.YYYY' -> ['Apr 25', ...] (Duplikat von constants.berechne_monate)."""
    monat_str, jahr_str = start_monat.split(".")
    monat_idx = int(monat_str) - 1
    jahr = int(jahr_str)
    monate = []
    for i in range(anzahl_monate):
        m = (monat_idx + i) % 12
        j = jahr + (monat_idx + i) // 12
        monate.append(f"{MONAT_NAMEN[m]} {j % 100:02d}")
    return monate


def read_config(wb) -> tuple[str, int]:
    ws = wb["Konfiguration"]
    start_monat = str(ws.cell(row=17, column=3).value).strip()
    anzahl_monate = int(ws.cell(row=18, column=3).value)
    return start_monat, anzahl_monate


def read_project_sheet(ws, anzahl_monate: int) -> dict:
    name = str(ws.cell(row=1, column=1).value or ws.title).strip()

    monate = [
        ws.cell(row=MONAT_HEADER_ROW, column=FIRST_MONAT_COL + i).value
        for i in range(anzahl_monate)
    ]

    subprojects = []
    for tp_idx, start_row in enumerate(TP_START_ROWS):
        header_row = start_row - 1
        tp_name = str(ws.cell(row=header_row, column=1).value or f"Teilprojekt {tp_idx + 1}").strip()

        phasen: dict[str, list[str]] = {}
        for ph_idx, code in enumerate(PHASE_ROW_CODES):
            row = start_row + ph_idx
            for i, monat in enumerate(monate):
                if not monat:
                    continue
                cell_val = ws.cell(row=row, column=FIRST_MONAT_COL + i).value
                if cell_val and str(cell_val).strip() == code:
                    phasen.setdefault(monat, []).append(code)

        fte: dict[str, float] = {}
        fte_row = TP_FTE_ROWS[tp_idx]
        for i, monat in enumerate(monate):
            if not monat:
                continue
            val = ws.cell(row=fte_row, column=FIRST_MONAT_COL + i).value
            if val:
                fte[monat] = float(val)

        # Leere Teilprojekte (keine Phasen, kein FTE, kein Name) überspringen.
        if tp_name.startswith("Teilprojekt ") and not phasen and not fte:
            continue

        subprojects.append({"name": tp_name, "reihenfolge": tp_idx, "phasen": phasen, "fte": fte})

    return {"name": name, "subprojects": subprojects}


def _verdichte_phasen(monate: list[str], phasen: dict[str, list[str]]) -> list[dict]:
    """Planphase-Erzeugung: aufeinanderfolgende Monate mit gleichem Code -> EINE Phase.

    Rückgabe: Liste von Dicts mit "kind": "plan_phase"|"milestone" und den zu
    schreibenden Feldern. '?'-Zellen werden zu Milestones.
    """
    idx_by_monat = {monat: i for i, monat in enumerate(monate)}

    runs: dict[str, list[int]] = {}
    for monat, codes in phasen.items():
        idx = idx_by_monat.get(monat)
        if idx is None:
            continue
        for code in codes:
            runs.setdefault(code, []).append(idx)

    result: list[dict] = []
    for code, indices in sorted(runs.items()):
        indices = sorted(set(indices))
        gruppen: list[list[int]] = []
        for idx in indices:
            if gruppen and idx == gruppen[-1][-1] + 1:
                gruppen[-1].append(idx)
            else:
                gruppen.append([idx])
        for gruppe in gruppen:
            start_label = monate[gruppe[0]]
            end_label = monate[gruppe[-1]]
            if code == "?":
                result.append(
                    {
                        "kind": "milestone",
                        "name": f"Meilenstein {_monat_end(end_label)}",
                        "forecast_date": _monat_end(end_label),
                    }
                )
            else:
                result.append(
                    {
                        "kind": "plan_phase",
                        "phase_type": PHASE_LABELS.get(code, f"Phase {code}"),
                        "forecast_start": _monat_start(start_label),
                        "forecast_end": _monat_end(end_label),
                    }
                )
    return result


def _ensure_default_role(api_base: str) -> int:
    """Findet die ResourceRole "Allgemein" und legt sie bei Bedarf an (identischer
    Default wie in Migration 0003_phase26_legacy_cutover). Rückgabe: role_id."""
    roles = requests.get(f"{api_base}/resource-roles", timeout=30)
    roles.raise_for_status()
    for role in roles.json():
        if role["name"] == "Allgemein":
            return role["id"]
    resp = requests.post(
        f"{api_base}/resource-roles",
        json={"name": "Allgemein", "description": "Default-Rolle für den Excel-Import (migration/import_excel.py)"},
        timeout=30,
    )
    resp.raise_for_status()
    return resp.json()["id"]


def push_to_api(api_base: str, start_monat: str, anzahl_monate: int, project: dict) -> None:
    resp = requests.post(
        f"{api_base}/projects",
        json={
            "name": project["name"],
            "start_monat": start_monat,
            "anzahl_monate": anzahl_monate,
        },
        timeout=30,
    )
    resp.raise_for_status()
    project_id = resp.json()["id"]

    monate = _berechne_monate(start_monat, anzahl_monate)
    role_id = _ensure_default_role(api_base)

    for sp in project["subprojects"]:
        resp = requests.post(
            f"{api_base}/projects/{project_id}/subprojects",
            json={"name": sp["name"], "reihenfolge": sp["reihenfolge"]},
            timeout=30,
        )
        resp.raise_for_status()
        subproject_id = resp.json()["id"]

        for eintrag in _verdichte_phasen(monate, sp["phasen"]):
            if eintrag["kind"] == "milestone":
                requests.post(
                    f"{api_base}/projects/{project_id}/milestones",
                    json={"subproject_id": subproject_id, **eintrag},
                    timeout=30,
                ).raise_for_status()
            else:
                requests.post(
                    f"{api_base}/projects/{project_id}/plan-phases",
                    json={"subproject_id": subproject_id, **eintrag},
                    timeout=30,
                ).raise_for_status()

        for monat, wert in sp["fte"].items():
            requests.post(
                f"{api_base}/projects/{project_id}/resource-demands",
                json={"resource_role_id": role_id, "period": monat, "fte": wert},
                timeout=30,
            ).raise_for_status()

    print(f"  -> importiert als Projekt #{project_id}: {project['name']}")


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("xlsm_path", help="Pfad zur bestehenden .xlsm-Datei")
    parser.add_argument("--api", default="http://localhost:8000", help="Basis-URL der Backend-API")
    parser.add_argument("--dry-run", action="store_true", help="Nur einlesen und ausgeben, nicht in die API schreiben")
    args = parser.parse_args()

    wb = load_workbook(args.xlsm_path, data_only=True)
    start_monat, anzahl_monate = read_config(wb)
    print(f"Startmonat: {start_monat}, Anzahl Monate: {anzahl_monate}")

    project_sheets = [ws for ws in wb.worksheets if ws.title not in SKIP_SHEETS]
    print(f"Gefundene Projektblätter: {len(project_sheets)}")

    for ws in project_sheets:
        project = read_project_sheet(ws, anzahl_monate)
        if args.dry_run:
            print(project)
        else:
            push_to_api(args.api, start_monat, anzahl_monate, project)

    if args.dry_run:
        print("Dry-run — es wurde nichts in die API geschrieben.")


if __name__ == "__main__":
    main()